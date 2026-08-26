# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).resolve().parents[1] / "presentacion-enlyce-resumida.pptx"

NAVY = RGBColor(0x12, 0x22, 0x33)
NAVY_2 = RGBColor(0x1B, 0x30, 0x46)
CREAM = RGBColor(0xF3, 0xEF, 0xE6)
PAPER = RGBColor(0xFA, 0xF7, 0xF0)
GOLD = RGBColor(0xB4, 0x91, 0x52)
INK = RGBColor(0x1F, 0x2A, 0x33)
SLATE = RGBColor(0x62, 0x6D, 0x75)
GREEN = RGBColor(0x35, 0x64, 0x51)
BURGUNDY = RGBColor(0x72, 0x2F, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_shape(slide, shape, x, y, w, h, fill, line=None, radius=None):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
        item.line.width = Pt(1)
    if radius is not None:
        try:
            item.adjustments[0] = radius
        except (IndexError, TypeError):
            pass
    return item


def add_text(slide, x, y, w, h, text, size=16, color=INK, bold=False,
             font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             margin=0, tracking=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if tracking is not None:
        run.font._element.set("spc", str(tracking))
    return box


def label(slide, x, y, w, text, fill=GOLD, color=NAVY):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill, radius=0.16)
    add_text(slide, x, y + 0.01, w, 0.28, text.upper(), 9, color, True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def base_slide(section, title, number):
    slide = prs.slides.add_slide(BLANK)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PAPER
    label(slide, 0.64, 0.48, 1.55, section)
    add_text(slide, 0.64, 0.98, 11.7, 0.62, title, 31, NAVY, True, "Georgia")
    add_text(slide, 12.25, 0.48, 0.42, 0.28, f"0{number}", 10, GOLD, True,
             align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title, body, accent=GOLD, fill=CREAM,
         title_size=14, body_size=11.5):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.08)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.07, h, accent)
    add_text(slide, x + 0.27, y + 0.19, w - 0.5, 0.34, title, title_size, NAVY, True)
    add_text(slide, x + 0.27, y + 0.62, w - 0.5, h - 0.78, body, body_size, SLATE)


# 01 - Portada
slide = prs.slides.add_slide(BLANK)
background = slide.background.fill
background.solid()
background.fore_color.rgb = NAVY
add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.16, 7.5, GOLD)

# Motivo arquitectonico sutil
add_shape(slide, MSO_SHAPE.ARC, 9.0, 0.7, 3.6, 3.6, NAVY, GOLD)
add_shape(slide, MSO_SHAPE.ARC, 9.7, 1.4, 2.2, 2.2, NAVY, GOLD)
add_shape(slide, MSO_SHAPE.RECTANGLE, 10.73, 3.05, 0.08, 3.4, GOLD)

label(slide, 0.85, 0.74, 1.75, "Actividad 1", GOLD, NAVY)
add_text(slide, 0.85, 1.55, 7.5, 1.0, "ENLYCE", 58, WHITE, True, "Georgia")
add_text(slide, 0.88, 2.65, 7.65, 0.75,
         "Modelo de negocio para un CRM inmobiliario colombiano", 24, CREAM)
add_text(slide, 0.88, 4.12, 6.9, 0.72,
         "Cada lead con responsable. Cada seguimiento visible.\nCada dato personal protegido.",
         17, WHITE, False)
add_text(slide, 0.88, 6.55, 7.2, 0.28,
         "Cliente ancla: L&C Propiedad Raíz  ·  Medellín  ·  Agosto de 2026",
         10.5, RGBColor(0xC7, 0xD0, 0xD8))


# 02 - Problema y propuesta
slide = base_slide("Oportunidad", "El problema no es captar leads: es no perderlos", 2)
add_text(slide, 0.64, 1.76, 5.1, 0.5,
         "L&C opera con cuatro asesores, menos de 50 inmuebles y herramientas fragmentadas.",
         14, SLATE)

sources = [
    ("WASI", "Un usuario compartido", NAVY_2),
    ("WHATSAPP", "El lead queda en un celular", GREEN),
    ("EXCEL", "Sin historial ni responsable", BURGUNDY),
]
for index, (title, body, color) in enumerate(sources):
    y = 2.52 + index * 1.06
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.64, y, 4.75, 0.82, CREAM, radius=0.08)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, y + 0.18, 1.14, 0.42, color, radius=0.14)
    add_text(slide, 0.82, y + 0.23, 1.14, 0.24, title, 9, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, 2.18, y + 0.21, 2.9, 0.34, body, 12.5, INK)

add_shape(slide, MSO_SHAPE.CHEVRON, 5.75, 3.18, 0.72, 1.1, GOLD)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.78, 1.78, 5.9, 3.45, NAVY, radius=0.05)
add_text(slide, 7.18, 2.12, 5.1, 0.34, "EFECTO EN EL NEGOCIO", 10, GOLD, True)
add_text(slide, 7.18, 2.7, 4.85, 1.45,
         "Leads olvidados\nAtención duplicada\nCero trazabilidad", 25, WHITE, True, "Georgia")
add_text(slide, 7.18, 4.5, 4.85, 0.34,
         "Y sin evidencia del consentimiento exigido por la Ley 1581.", 12, CREAM)

add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.64, 5.68, 12.04, 1.08, GREEN, radius=0.06)
add_text(slide, 0.97, 5.89, 2.15, 0.2, "PROPUESTA DE VALOR", 9, GOLD, True)
add_text(slide, 3.08, 5.82, 9.12, 0.43,
         "Enlyce ordena el negocio: responsable único, pipeline visible, alertas y cumplimiento integrado.",
         16, WHITE, True)


# 03 - Canvas resumido
slide = base_slide("Canvas", "Un modelo enfocado en micro-inmobiliarias", 3)
card(slide, 0.64, 1.78, 3.72, 2.03, "Cliente ideal",
     "Inmobiliarias independientes de Medellín; 1–8 asesores; menos de 50 inmuebles; hoy usan WhatsApp, Excel o un CRM limitado.",
     GREEN, CREAM, 15, 11.5)
card(slide, 4.56, 1.78, 3.72, 2.03, "Valor entregado",
     "Multi-asesor nativo, pipeline Kanban, seguimiento automático, implantación guiada y cumplimiento de la Ley 1581.",
     GOLD, CREAM, 15, 11.5)
card(slide, 8.48, 1.78, 4.2, 2.03, "Relación y canales",
     "Aplicación web, demostraciones, soporte directo por WhatsApp y correo, capacitación y crecimiento por referidos.",
     NAVY_2, CREAM, 15, 11.5)

add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.64, 4.08, 5.85, 2.07, NAVY, radius=0.05)
add_text(slide, 0.97, 4.4, 2.5, 0.28, "INGRESOS", 10, GOLD, True)
add_text(slide, 0.97, 4.88, 4.9, 0.46, "$89.000 base + $19.000", 23, WHITE, True, "Georgia")
add_text(slide, 0.97, 5.43, 4.95, 0.35,
         "por asesor adicional · setup único de $150.000 · complementos opcionales", 11, CREAM)

add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.78, 4.08, 5.9, 2.07, CREAM, radius=0.05)
add_text(slide, 7.11, 4.4, 2.5, 0.28, "OPERACIÓN", 10, GREEN, True)
add_text(slide, 7.11, 4.86, 4.95, 0.42, "Plataforma propia + servicio cercano", 18, NAVY, True, "Georgia")
add_text(slide, 7.11, 5.4, 4.95, 0.42,
         "Costos: hosting, PostgreSQL, mensajería, correo, monitoreo y soporte.", 11.5, SLATE)

add_text(slide, 0.64, 6.48, 12.0, 0.3,
         "Foco deliberado: resolver muy bien el tramo entre el CRM gratuito insuficiente y la suite costosa.",
         11.5, BURGUNDY, True, align=PP_ALIGN.CENTER)


# 04 - Modelado del negocio
slide = base_slide("Modelado", "Del primer contacto al cierre, con reglas explícitas", 4)

stages = ["REGISTRO", "CONSENTIMIENTO", "ASIGNACIÓN", "SEGUIMIENTO", "CIERRE"]
colors = [NAVY_2, BURGUNDY, GOLD, GREEN, NAVY_2]
for index, (stage, color) in enumerate(zip(stages, colors)):
    x = 0.64 + index * 2.43
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.78, 2.06, 0.86, color, radius=0.06)
    add_text(slide, x + 0.08, 2.03, 1.9, 0.25, stage, 9.5, WHITE, True, align=PP_ALIGN.CENTER)
    if index < len(stages) - 1:
        add_shape(slide, MSO_SHAPE.CHEVRON, x + 2.09, 2.03, 0.28, 0.34, GOLD)

model_cards = [
    ("ACTORES", "Asesor · administrador · lead · propietario", GREEN),
    ("OBJETOS", "Lead · inmueble · interacción · visita · consentimiento", GOLD),
    ("REGLAS", "Un responsable · avance controlado · baja lógica · auditoría", BURGUNDY),
]
for index, (title, body, accent) in enumerate(model_cards):
    card(slide, 0.64 + index * 4.02, 3.08, 3.78, 1.62, title, body, accent, CREAM, 13, 11.2)

add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.64, 5.08, 12.04, 1.31, NAVY, radius=0.05)
add_text(slide, 0.96, 5.34, 2.05, 0.26, "REGLA TRANSVERSAL", 9.5, GOLD, True)
add_text(slide, 3.05, 5.23, 9.12, 0.55,
         "Ningún lead se crea sin autorización previa, expresa e informada; toda operación queda auditada.",
         15.5, WHITE, True)
add_text(slide, 0.64, 6.62, 12.0, 0.26,
         "Los procesos se vuelven casos de uso; las reglas viven en el dominio; los objetos se persisten con trazabilidad.",
         10.5, SLATE, align=PP_ALIGN.CENTER)


# 05 - Viabilidad y cierre
slide = prs.slides.add_slide(BLANK)
background = slide.background.fill
background.solid()
background.fore_color.rgb = NAVY
label(slide, 0.72, 0.55, 1.4, "Cierre", GOLD, NAVY)
add_text(slide, 0.72, 1.12, 7.0, 0.88,
         "Un negocio pequeño no necesita un CRM pequeño.", 32, WHITE, True, "Georgia")
add_text(slide, 0.72, 2.13, 6.85, 0.62,
         "Necesita uno enfocado: simple de adoptar, difícil de ignorar y preparado para cumplir.",
         16, CREAM)

metrics = [
    ("172", "pruebas automatizadas\nen verde"),
    ("4", "capas de arquitectura\nlimpia"),
    ("100 %", "consentimiento exigido\npor diseño"),
]
for index, (value, caption) in enumerate(metrics):
    x = 0.72 + index * 2.42
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.28, 2.14, 1.6, NAVY_2, GOLD, 0.06)
    add_text(slide, x, 3.52, 2.14, 0.52, value, 28, GOLD, True, "Georgia", PP_ALIGN.CENTER)
    add_text(slide, x + 0.12, 4.17, 1.9, 0.42, caption, 9.5, CREAM, False,
             align=PP_ALIGN.CENTER)

add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.45, 1.15, 4.16, 4.97, CREAM, radius=0.04)
add_text(slide, 8.83, 1.56, 3.35, 0.3, "SIGUIENTE VALIDACIÓN", 10, GREEN, True)
add_text(slide, 8.83, 2.08, 3.18, 0.52, "L&C operando con datos reales", 18, NAVY, True, "Georgia")

validation = [
    ("01", "50+ leads registrados"),
    ("02", "Primera respuesta < 2 horas"),
    ("03", "100 % con consentimiento"),
    ("04", "Disponibilidad > 99 %"),
]
for index, (number, text) in enumerate(validation):
    y = 2.91 + index * 0.63
    add_text(slide, 8.83, y, 0.42, 0.26, number, 9, GOLD, True)
    add_text(slide, 9.35, y - 0.03, 2.75, 0.32, text, 11.5, INK)

add_text(slide, 0.72, 6.35, 7.1, 0.4,
         "Enlyce convierte desorden comercial en procesos, responsables y evidencia.",
         14.5, WHITE, True)
add_text(slide, 0.72, 6.88, 7.1, 0.22, "Actividad 1 · Modelo de negocio", 9.5, GOLD)


prs.save(OUTPUT)
print(f"Creado: {OUTPUT}")
print(f"Diapositivas: {len(prs.slides)}")
