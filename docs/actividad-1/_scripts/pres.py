# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from acentos import fix

OUT = r"D:/Proyectos/Enlyce/docs/actividad-1/presentacion-enlyce.pptx"

NAVY = RGBColor(0x0F, 0x25, 0x40)
BLUE = RGBColor(0x1F, 0x6F, 0xEB)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
GREY = RGBColor(0x5B, 0x6B, 0x7F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1B, 0x8A, 0x5A)
AMBER = RGBColor(0xB4, 0x7A, 0x0A)
RED = RGBColor(0xB3, 0x2D, 0x2D)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = 13.333
BLANK = prs.slide_layouts[6]


def txt(sl, x, y, w, h, text, size=16, bold=False, color=NAVY, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font="Calibri", space=6, line=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    lines = text if isinstance(text, list) else text.splitlines()
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if line:
            p.line_spacing = line
        r = p.add_run()
        r.text = fix(ln)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def rect(sl, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    s.text_frame.text = ""
    return s


def slide(titulo, kicker=None):
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, W, 0.16, BLUE, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        txt(sl, 0.7, 0.42, 11, 0.3, kicker.upper(), size=11, bold=True, color=BLUE)
        txt(sl, 0.7, 0.72, 12, 0.6, titulo, size=30, bold=True, color=NAVY)
    else:
        txt(sl, 0.7, 0.55, 12, 0.8, titulo, size=30, bold=True, color=NAVY)
    txt(sl, 0.7, 6.95, 12, 0.3, "Enlyce - CRM Inmobiliario  |  L&C Propiedad Raiz", size=9, color=GREY)
    return sl


def bullets(sl, x, y, w, items, size=15, gap=0.52, color=NAVY):
    for i, it in enumerate(items):
        yy = y + i * gap
        d = rect(sl, x, yy + 0.10, 0.12, 0.12, BLUE, shape=MSO_SHAPE.OVAL)
        txt(sl, x + 0.32, yy, w, gap, it, size=size, color=color)


def card(sl, x, y, w, h, titulo, cuerpo, acento=BLUE, fondo=LIGHT, ts=15, cs=12):
    rect(sl, x, y, w, h, fondo)
    rect(sl, x, y, 0.07, h, acento, shape=MSO_SHAPE.RECTANGLE)
    txt(sl, x + 0.28, y + 0.18, w - 0.5, 0.4, titulo, size=ts, bold=True, color=NAVY)
    txt(sl, x + 0.28, y + 0.62, w - 0.5, h - 0.75, cuerpo, size=cs, color=GREY, space=3)


def tabla(sl, x, y, w, h, headers, rows, fs=11, hs=11, widths=None):
    shp = sl.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    t = shp.table
    if widths:
        total = sum(widths)
        for i, ww in enumerate(widths):
            t.columns[i].width = Emu(int(Inches(w) * ww / total))
    for i, hd in enumerate(headers):
        c = t.cell(0, i)
        c.text = fix(hd)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        p = c.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(hs); p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = WHITE; p.runs[0].font.name = "Calibri"
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for r, row in enumerate(rows, start=1):
        for i, v in enumerate(row):
            c = t.cell(r, i)
            c.text = fix(str(v))
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            p = c.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(fs)
            p.runs[0].font.color.rgb = NAVY
            p.runs[0].font.name = "Calibri"
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return t


# ---------------------------------------------------------------- 1 portada
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, 7.5, NAVY, shape=MSO_SHAPE.RECTANGLE)
rect(sl, 0, 0, 0.35, 7.5, BLUE, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 1.2, 1.9, 11, 0.4, "MODELO DE NEGOCIO  -  ACTIVIDAD 1", size=14, bold=True, color=RGBColor(0x7F, 0xB3, 0xFF))
txt(sl, 1.2, 2.4, 11, 1.2, "ENLYCE", size=66, bold=True, color=WHITE)
txt(sl, 1.2, 3.6, 11, 0.6, "CRM inmobiliario para micro-inmobiliarias colombianas", size=24, color=RGBColor(0xD8, 0xE4, 0xF2))
rect(sl, 1.2, 4.45, 3.0, 0.05, BLUE, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 1.2, 4.8, 11, 1.2,
    ["Cliente ancla: L&C Propiedad Raiz  -  Medellin",
     "Ingenieria de Software  -  Agosto de 2026",
     "ASP.NET Core 10 - PostgreSQL - React  |  172 pruebas automatizadas en verde"],
    size=14, color=RGBColor(0xB9, 0xCB, 0xE2), space=5)

# ---------------------------------------------------------------- 2 problema
sl = slide("Una inmobiliaria que pierde clientes por desorden, no por falta de demanda", "El problema")
txt(sl, 0.7, 1.75, 12, 0.5,
    "L&C Propiedad Raiz: 4 asesores, menos de 50 inmuebles activos, cero herramientas de gestion comercial.",
    size=15, color=GREY)
datos = [("Wasi gratis", "1 solo usuario para todo el equipo"),
         ("WhatsApp", "El lead vive en el celular de quien lo contesto"),
         ("Excel", "Sin historial, sin etapas, sin responsable"),
         ("Ley 1581", "Cero evidencia de autorizacion de datos")]
for i, (a, b) in enumerate(datos):
    card(sl, 0.7 + i * 3.12, 2.5, 2.9, 1.5, a, b, acento=RED, ts=16, cs=12)
rect(sl, 0.7, 4.4, 11.9, 1.9, LIGHT)
txt(sl, 1.1, 4.65, 11.1, 0.4, "Consecuencia medible", size=16, bold=True, color=NAVY)
bullets(sl, 1.1, 5.1, 10.8,
        ["Leads que nadie atiende porque nadie es el responsable formal.",
         "Dos asesores llamando al mismo cliente: mala experiencia y conflicto interno.",
         "Ninguna cifra de productividad: no se sabe quien convierte ni cuanto tarda."],
        size=13, gap=0.38)

# ---------------------------------------------------------------- 3 solucion
sl = slide("Enlyce ordena el negocio, no solo lo digitaliza", "La propuesta de valor")
txt(sl, 0.7, 1.7, 12, 0.5,
    "\"El CRM con el que una micro-inmobiliaria colombiana deja de perder leads.\"",
    size=19, bold=True, color=BLUE)
items = [("Multi-asesor desde el dia 1", "Cada lead tiene un unico responsable. Nadie se pisa el cliente."),
         ("Pipeline Kanban visible", "Contacto, Visita, Propuesta, Negociacion y Cierre en un tablero."),
         ("Alertas de seguimiento", "El lead olvidado se convierte en una tarea visible, no en una perdida."),
         ("Ley 1581 integrada", "Consentimiento obligatorio, politica versionada y auditoria completa."),
         ("Operativo en menos de 24 h", "Implantacion, migracion y capacitacion el mismo dia."),
         ("Precio intermedio", "~$158.000/mes: ni gratis sin valor ni costoso e inaccesible.")]
for i, (a, b) in enumerate(items):
    card(sl, 0.7 + (i % 3) * 4.0, 2.5 + (i // 3) * 1.95, 3.8, 1.75, a, b, ts=15, cs=12)

# ---------------------------------------------------------------- 4 mapa del modelado
sl = slide("Los cinco componentes del modelado", "Modelado del negocio")
comp = [("1. PROCESOS", "Secuencias de\nactividades que\ngeneran valor", BLUE),
        ("2. ACTIVIDADES", "Tareas concretas\ndentro de cada\nproceso", BLUE),
        ("3. REGLAS", "Normas que el\nsistema obliga\na cumplir", RED),
        ("4. ACTORES", "Personas y sistemas\nque interactuan\ncon el negocio", GREEN),
        ("5. OBJETOS", "Informacion que\ncircula por los\nprocesos", AMBER)]
for i, (a, b, col) in enumerate(comp):
    x = 0.7 + i * 2.44
    rect(sl, x, 2.3, 2.24, 2.6, LIGHT)
    rect(sl, x, 2.3, 2.24, 0.5, col, shape=MSO_SHAPE.RECTANGLE)
    txt(sl, x, 2.38, 2.24, 0.4, a, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(sl, x + 0.2, 3.05, 1.9, 1.7, b.split("\n"), size=12, color=GREY, align=PP_ALIGN.CENTER, space=1)
txt(sl, 0.7, 5.3, 11.9, 1.0,
    "En Enlyce estos cinco componentes no se quedan en el papel: cada uno tiene su reflejo directo en el codigo. "
    "Los procesos son casos de uso, las reglas viven en el dominio, los actores son roles autenticados y los objetos "
    "son entidades persistidas con auditoria.",
    size=14, color=NAVY)

# ---------------------------------------------------------------- 5 procesos
sl = slide("Proceso central: del primer contacto al cierre", "1. Procesos de negocio")
etapas = ["Registro\ndel lead", "Consentimiento\nLey 1581", "Asignacion\na un asesor", "Contacto e\ninteracciones",
          "Visita al\ninmueble", "Propuesta y\nnegociacion", "Cierre\n(ganado/perdido)"]
for i, e in enumerate(etapas):
    x = 0.7 + i * 1.72
    col = GREEN if i == len(etapas) - 1 else (RED if i == 1 else BLUE)
    rect(sl, x, 2.2, 1.55, 1.3, col)
    txt(sl, x + 0.08, 2.42, 1.4, 1.0, e.split("\n"), size=11, bold=True, color=WHITE,
        align=PP_ALIGN.CENTER, space=1)
    if i < len(etapas) - 1:
        txt(sl, x + 1.5, 2.65, 0.3, 0.4, ">", size=18, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
txt(sl, 0.7, 3.75, 11.9, 0.4, "Procesos de apoyo", size=17, bold=True, color=NAVY)
tabla(sl, 0.7, 4.2, 11.9, 2.2,
      ["Proceso de apoyo", "Que hace", "Valor que genera"],
      [["Gestion de inmuebles", "Alta, actualizacion y baja logica con propietario asociado", "Catalogo unico y confiable"],
       ["Control de seguimiento", "Detecta leads sin gestion y genera alertas", "Ningun lead se pierde por olvido"],
       ["Cumplimiento Ley 1581", "Consentimiento, politica, derechos del titular y auditoria", "Evita sanciones de la SIC"],
       ["Administracion de usuarios", "Asesores, autenticacion y control de roles", "Cada quien ve lo que le toca"]],
      widths=[2.6, 5.2, 3.2])

# ---------------------------------------------------------------- 6 actividades
sl = slide("Cada modulo del sistema es un grupo de actividades", "2. Actividades del negocio")
acts = [("Leads", "Registrar, capturar consentimiento, consultar, actualizar, dar de baja logica"),
        ("Pipeline", "Consultar tablero, mover de etapa, registrar transicion, ver historial"),
        ("Asignacion", "Asignar a un asesor, reasignar, consultar carga de trabajo"),
        ("Interacciones", "Registrar llamada, mensaje, correo o nota; ver historial del lead"),
        ("Visitas", "Agendar, confirmar, reprogramar, marcar realizada o cancelada"),
        ("Alertas", "Detectar leads sin seguimiento, notificar y marcar como atendida"),
        ("Inmuebles", "Publicar, asociar propietario, actualizar precio y estado, retirar"),
        ("Datos personales", "Otorgar y revocar consentimiento, publicar politica, atender al titular"),
        ("Seguridad", "Iniciar y cerrar sesion, validar rol, registrar el evento en auditoria")]
for i, (a, b) in enumerate(acts):
    card(sl, 0.7 + (i % 3) * 4.0, 1.9 + (i // 3) * 1.6, 3.8, 1.4, a, b, ts=14, cs=11)
txt(sl, 0.7, 6.75, 11.9, 0.3, "9 modulos implementados y expuestos como endpoints de la API.",
    size=12, bold=True, color=BLUE)

# ---------------------------------------------------------------- 7 reglas
sl = slide("Las reglas no son un manual: el sistema las obliga", "3. Reglas del negocio")
rect(sl, 0.7, 1.85, 3.8, 4.6, RGBColor(0xFD, 0xEC, 0xEC))
rect(sl, 0.7, 1.85, 3.8, 0.5, RED, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 0.95, 1.94, 3.4, 0.4, "LEGALES (LEY 1581)", size=13, bold=True, color=WHITE)
txt(sl, 0.95, 2.55, 3.3, 3.7,
    ["- Ningun lead se crea sin autorizacion previa, expresa e informada.",
     "- La autorizacion guarda fecha, medio, finalidad y version de politica.",
     "- El titular puede revocar en cualquier momento.",
     "- Derechos de conocer, actualizar, rectificar y suprimir.",
     "- La politica es publica, versionada y debe estar vigente.",
     "- Toda operacion sobre datos queda auditada."], size=11.5, color=NAVY, space=7)

rect(sl, 4.75, 1.85, 3.8, 4.6, RGBColor(0xE8, 0xF1, 0xFF))
rect(sl, 4.75, 1.85, 3.8, 0.5, BLUE, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 5.0, 1.94, 3.4, 0.4, "COMERCIALES", size=13, bold=True, color=WHITE)
txt(sl, 5.0, 2.55, 3.3, 3.7,
    ["- Todo lead tiene exactamente un asesor responsable.",
     "- Un lead solo avanza por las etapas definidas, sin saltarse ninguna.",
     "- Un lead perdido exige motivo de descarte registrado.",
     "- Un lead sin interaccion en el plazo definido genera alerta automatica.",
     "- Solo se agenda visita sobre inmueble disponible y lead activo."], size=11.5, color=NAVY, space=8)

rect(sl, 8.8, 1.85, 3.8, 4.6, RGBColor(0xE8, 0xF7, 0xEC))
rect(sl, 8.8, 1.85, 3.8, 0.5, GREEN, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 9.05, 1.94, 3.4, 0.4, "INTEGRIDAD", size=13, bold=True, color=WHITE)
txt(sl, 9.05, 2.55, 3.3, 3.7,
    ["- Nada se elimina fisicamente: toda baja es logica y auditada.",
     "- El correo del asesor es unico en el sistema.",
     "- Las contrasenas se guardan con hash bcrypt, nunca en texto plano.",
     "- Cada peticion se autentica con token en cookie HttpOnly y se autoriza por rol.",
     "- Un inmueble con visitas pendientes no se puede retirar."], size=11.5, color=NAVY, space=8)

# ---------------------------------------------------------------- 8 actores
sl = slide("Quien interactua con el negocio", "4. Actores del negocio")
tabla(sl, 0.7, 1.85, 11.9, 4.6,
      ["Actor", "Tipo", "Rol en el negocio"],
      [["Asesor comercial", "Humano - interno", "Registra y atiende leads, agenda visitas, mueve el pipeline. Usuario principal."],
       ["Administrador o gerente", "Humano - interno", "Gestiona asesores, revisa productividad, configura etapas y consulta auditoria."],
       ["Lead o cliente interesado", "Humano - externo", "Titular de los datos. Pide informacion, visita inmuebles y compra o arrienda."],
       ["Propietario del inmueble", "Humano - externo", "Entrega el inmueble en gestion a la inmobiliaria."],
       ["Oficial de datos personales", "Humano - interno", "Atiende solicitudes del titular y vela por la Ley 1581."],
       ["Superintendencia de Industria y Comercio", "Organizacion - externa", "Vigila el tratamiento de datos y exige el registro en el RNBD."],
       ["WhatsApp Business API (Meta)", "Sistema - externo", "Canal de entrada de leads y de comunicacion."],
       ["Motor de alertas de Enlyce", "Sistema - interno", "Detecta leads sin seguimiento y notifica sin intervencion humana."]],
      fs=11, widths=[3.0, 2.4, 6.5])

# ---------------------------------------------------------------- 9 objetos
sl = slide("La informacion que circula por el negocio", "5. Objetos del negocio")
objs = [("Lead", "Nombre, contacto, fuente, etapa, asesor, estado"),
        ("Asesor", "Nombre, correo, hash de contrasena, rol, activo"),
        ("Inmueble", "Codigo, tipo, direccion, precio, area, estado"),
        ("Propietario", "Nombre, documento, contacto"),
        ("Interaccion", "Lead, tipo, fecha, descripcion, asesor"),
        ("Visita", "Lead, inmueble, fecha, estado, observaciones"),
        ("Etapa de pipeline", "Nombre, orden, condiciones de avance"),
        ("Transicion", "Lead, etapa origen y destino, fecha, asesor"),
        ("Consentimiento", "Titular, finalidad, medio, fecha, version, estado"),
        ("Politica de tratamiento", "Version, contenido, fecha de vigencia"),
        ("Registro de auditoria", "Usuario, accion, entidad, fecha, valores"),
        ("Alerta de seguimiento", "Lead, asesor, motivo, fecha, estado")]
for i, (a, b) in enumerate(objs):
    card(sl, 0.7 + (i % 4) * 3.0, 1.9 + (i // 4) * 1.55, 2.85, 1.35, a, b,
         acento=AMBER, ts=13, cs=10.5)
txt(sl, 0.7, 6.6, 11.9, 0.4,
    "Cada objeto es una entidad del dominio, sin dependencias externas y cubierta por pruebas unitarias.",
    size=12, bold=True, color=BLUE)

# ---------------------------------------------------------------- 10 canvas
sl = slide("Business Model Canvas de Enlyce", "El modelo de negocio en una hoja")
bloques = [
    ("SOCIOS CLAVE", "LYC (cliente ancla), Azure/.NET, Meta WhatsApp API, correo transaccional, universidad, comunidad open source", 0.7, 1.85, 2.35, 2.15, BLUE),
    ("ACTIVIDADES CLAVE", "Desarrollo del CRM, implantacion en 24 h, capacitacion, soporte, cumplimiento Ley 1581", 3.15, 1.85, 2.35, 1.02, BLUE),
    ("RECURSOS CLAVE", "Plataforma propia, arquitectura limpia, modulo de cumplimiento, datos de LYC, equipo", 3.15, 2.98, 2.35, 1.02, BLUE),
    ("PROPUESTA DE VALOR", "Multi-asesor real, pipeline Kanban, alertas de seguimiento, Ley 1581 integrada, listo en 24 h, precio intermedio", 5.6, 1.85, 2.35, 2.15, GREEN),
    ("RELACION CON CLIENTES", "Implantacion acompanada, soporte por WhatsApp, co-creacion con el cliente ancla", 8.05, 1.85, 2.35, 1.02, BLUE),
    ("CANALES", "App web, WhatsApp Business, correo, demostracion, referidos del gremio", 8.05, 2.98, 2.35, 1.02, BLUE),
    ("SEGMENTOS DE CLIENTES", "Inmobiliarias de 1 a 8 asesores, menos de 50 inmuebles, Medellin y area metropolitana", 10.5, 1.85, 2.1, 2.15, BLUE),
    ("ESTRUCTURA DE COSTOS", "Hosting $120k-$200k/mes, dominio, WhatsApp por conversacion, correo y monitoreo, horas de soporte", 0.7, 4.15, 5.8, 1.55, AMBER),
    ("FUENTES DE INGRESO", "Suscripcion $89k + $19k por asesor extra, implantacion $150k, complementos de WhatsApp, portal y reportes", 6.8, 4.15, 5.8, 1.55, GREEN),
]
for tit, cue, x, y, w, h, col in bloques:
    rect(sl, x, y, w, h, LIGHT)
    rect(sl, x, y, w, 0.32, col, shape=MSO_SHAPE.RECTANGLE)
    txt(sl, x + 0.12, y + 0.03, w - 0.2, 0.3, tit, size=9.5, bold=True, color=WHITE)
    txt(sl, x + 0.12, y + 0.42, w - 0.24, h - 0.5, cue, size=9.5, color=NAVY, space=0)
txt(sl, 0.7, 5.9, 11.9, 0.5,
    "Restriccion transversal: la Ley 1581 de 2012 exige autorizacion previa, expresa e informada. "
    "Ningun lead se crea sin consentimiento registrado.", size=12, bold=True, color=RED)

# ---------------------------------------------------------------- 11 segmento
sl = slide("A quien le sirve Enlyce (y a quien no)", "Segmento de clientes")
card(sl, 0.7, 1.85, 5.8, 2.5, "Cliente ideal",
     "- Inmobiliaria independiente de Medellin y su area metropolitana.\n"
     "- De 1 a 8 asesores comerciales.\n"
     "- Menos de 50 inmuebles activos.\n"
     "- Hoy usa Wasi gratis, Excel y WhatsApp sin orden.\n"
     "- Factura entre $100 M y $2.000 M COP al ano.", acento=GREEN, ts=17, cs=13)
card(sl, 6.8, 1.85, 5.8, 2.5, "No es cliente",
     "- Inmobiliarias de mas de 15 asesores: les sirve Tokko.\n"
     "- Constructoras grandes con ERP propio.\n"
     "- Portales de clasificados.\n"
     "- Quien no quiere cambiar de herramienta ni por acompanamiento.",
     acento=RED, fondo=RGBColor(0xFD, 0xEC, 0xEC), ts=17, cs=13)
txt(sl, 0.7, 4.6, 11.9, 0.4, "Segmentos secundarios", size=17, bold=True, color=NAVY)
bullets(sl, 0.7, 5.05, 11.5,
        ["Asesores independientes que estan a punto de formar equipo.",
         "Franquicias locales de inmobiliarias grandes que necesitan autonomia.",
         "Constructoras pequenas con venta directa de unidades nuevas."], size=14, gap=0.42)

# ---------------------------------------------------------------- 12 ingresos
sl = slide("Como gana dinero Enlyce", "Fuentes de ingreso")
tabla(sl, 0.7, 1.8, 7.2, 2.8,
      ["Concepto", "Precio (COP)"],
      [["Suscripcion base por equipo (hasta 3 asesores)", "$89.000 / mes"],
       ["Asesor adicional (maximo 8)", "$19.000 / mes"],
       ["Implantacion inicial (migracion + capacitacion)", "$150.000 unico"],
       ["Complemento WhatsApp Business API", "$45.000 / mes"],
       ["Complemento portal web publico", "$35.000 / mes"],
       ["Complemento de reportes avanzados", "$25.000 / mes"]],
      fs=12, widths=[5.2, 2.0])
card(sl, 8.2, 1.8, 4.4, 2.8, "Caso tipico: LYC",
     "4 asesores\n\n"
     "Suscripcion base      $89.000\n"
     "1 asesor adicional    $19.000\n"
     "Implantacion (3 meses) $50.000\n"
     "-------------------------------\n"
     "Total mensual        $158.000", acento=GREEN, ts=17, cs=13)
txt(sl, 0.7, 4.85, 11.9, 0.4, "Proyeccion del primer semestre", size=17, bold=True, color=NAVY)
tabla(sl, 0.7, 5.3, 11.9, 1.1,
      ["Mes", "1", "2", "3", "4", "5", "6"],
      [["Clientes", "1", "2", "3", "5", "7", "10"],
       ["MRR (COP)", "158k", "316k", "474k", "790k", "1.106k", "1.580k"]],
      fs=12)

# ---------------------------------------------------------------- 13 costos
sl = slide("Estructura de costos y punto de equilibrio", "Costos")
tabla(sl, 0.7, 1.85, 7.2, 3.0,
      ["Categoria", "Costo mensual (COP)"],
      [["Hosting Azure + PostgreSQL administrado", "$120.000 - $200.000"],
       ["Dominio y certificados (SSL gratuito)", "~$4.000"],
       ["WhatsApp Business API", "Variable por conversacion"],
       ["Correo transaccional y monitoreo", "$0 - $35.000"],
       ["Desarrollo (96 h/semestre)", "Costo hundido academico"],
       ["Total fijo estimado", "$125.000 - $240.000"]],
      fs=12, widths=[4.8, 2.4])
card(sl, 8.2, 1.85, 4.4, 3.0, "Punto de equilibrio",
     "Con un solo cliente pagando $158.000/mes el modelo ya cubre el piso de costos.\n\n"
     "Desde el segundo cliente el margen es positivo, porque la infraestructura es compartida "
     "y el costo marginal por cliente adicional tiende a cero.\n\n"
     "Es un modelo dirigido por el valor, no por el costo.", acento=AMBER,
     fondo=RGBColor(0xFF, 0xF6, 0xE5), ts=17, cs=13)
txt(sl, 0.7, 5.2, 11.9, 1.0,
    "Riesgo asumido: el costo real de hosting solo se confirma en produccion. La estimacion se basa en precios "
    "de lista de Azure y no ha sido validada con trafico real; se marcara como comprobada al desplegar.",
    size=12, color=GREY)

# ---------------------------------------------------------------- 14 competencia
sl = slide("Por que no basta con lo que ya existe", "Competencia")
tabla(sl, 0.7, 1.85, 11.9, 3.6,
      ["Factor", "Enlyce", "Wasi", "Tokko", "Witei"],
      [["Multi-asesor real", "Si, es el nucleo", "No (1 usuario)", "Si, pero costoso", "Si"],
       ["Pipeline Kanban", "Si", "No", "Si", "Si"],
       ["Cumplimiento Ley 1581", "Integrado", "No", "Parcial", "No"],
       ["Alertas de seguimiento", "Si, automaticas", "No", "Si", "Parcial"],
       ["Tiempo de implantacion", "Menos de 24 h", "Inmediato", "1 a 2 semanas", "Dias"],
       ["Precio (4 asesores)", "~$158.000/mes", "Gratis limitado", "$313k - $667k/mes", "EUR 25-65/mes"],
       ["Soporte local en Colombia", "Si", "Si", "Si", "No"]],
      fs=12, widths=[3.2, 2.4, 2.1, 2.2, 2.0])
txt(sl, 0.7, 5.7, 11.9, 0.8,
    "El hueco del mercado esta entre el gratuito que no sirve para un equipo y el caro que una micro-inmobiliaria "
    "no puede pagar. Enlyce ocupa exactamente ese espacio, y ademas resuelve el cumplimiento normativo que nadie cubre.",
    size=14, color=NAVY)

# ---------------------------------------------------------------- 15 estado
sl = slide("El modelo ya esta construido, no solo planteado", "Estado actual del producto")
kpis = [("172", "pruebas automatizadas\nen verde"), ("9", "modulos funcionales\nen la API"),
        ("4", "capas de arquitectura\nlimpia"), ("100 %", "de leads con\nconsentimiento")]
for i, (n, d) in enumerate(kpis):
    x = 0.7 + i * 3.05
    rect(sl, x, 1.85, 2.85, 1.5, NAVY)
    txt(sl, x, 1.98, 2.85, 0.6, n, size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(sl, x, 2.62, 2.85, 0.7, d.split("\n"), size=11, color=RGBColor(0xB9, 0xCB, 0xE2),
        align=PP_ALIGN.CENTER, space=0)
txt(sl, 0.7, 3.6, 11.9, 0.4, "Que esta implementado", size=17, bold=True, color=NAVY)
bullets(sl, 0.7, 4.05, 5.6,
        ["Leads, inmuebles y propietarios con baja logica.",
         "Pipeline con etapas, transiciones y asignacion.",
         "Interacciones y visitas por lead.",
         "Alertas de seguimiento automaticas."], size=13, gap=0.42)
bullets(sl, 6.8, 4.05, 5.6,
        ["Autenticacion JWT con roles y cookie HttpOnly.",
         "Contrasenas con hash bcrypt.",
         "Consentimiento, politica y auditoria (Ley 1581).",
         "Frontend React: login, dashboard, pipeline, leads."], size=13, gap=0.42)

# ---------------------------------------------------------------- 16 roadmap
sl = slide("Lo que sigue", "Roadmap del semestre")
fases = [("Fase 1  -  Fundacion", "Dominio, arquitectura limpia, PostgreSQL, CRUD de leads e inmuebles, pruebas de dominio.", GREEN, "COMPLETADA"),
         ("Fase 2  -  Pipeline", "Kanban, asignacion por asesor, interacciones, visitas, alertas y pruebas de integracion.", GREEN, "COMPLETADA"),
         ("Fase 3  -  Cumplimiento", "Consentimiento, politica versionada, derechos del titular, auditoria y frontend React.", AMBER, "EN CURSO"),
         ("Fase 4  -  Entrega", "Reportes, despliegue en la nube, documentacion y demostracion con datos reales de LYC.", GREY, "PENDIENTE")]
for i, (a, b, col, est) in enumerate(fases):
    y = 1.9 + i * 1.2
    rect(sl, 0.7, y, 11.9, 1.05, LIGHT)
    rect(sl, 0.7, y, 0.09, 1.05, col, shape=MSO_SHAPE.RECTANGLE)
    txt(sl, 1.0, y + 0.14, 6.5, 0.35, a, size=16, bold=True, color=NAVY)
    txt(sl, 1.0, y + 0.55, 8.8, 0.4, b, size=12, color=GREY)
    rect(sl, 10.7, y + 0.3, 1.6, 0.45, col)
    txt(sl, 10.7, y + 0.38, 1.6, 0.3, est, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, 0.7, 6.75, 11.9, 0.3,
    "Fuera de alcance este semestre: pasarela de pagos, contratos digitales, app movil nativa e integracion con portales.",
    size=12, color=GREY)

# ---------------------------------------------------------------- 17 metricas
sl = slide("Como sabremos que el modelo funciona", "Metricas de exito")
tabla(sl, 0.7, 1.9, 11.9, 3.4,
      ["Metrica", "Meta del semestre", "Por que importa"],
      [["LYC operando en Enlyce", "Cliente ancla activo", "Valida que el modelo resuelve un dolor real"],
       ["Leads registrados", "Mas de 50", "Prueba que el sistema se usa a diario"],
       ["Tiempo de primera respuesta", "Menos de 2 horas", "Es el factor que mas influye en la conversion"],
       ["Leads con consentimiento", "100 %", "Cumplimiento de la Ley 1581 sin excepciones"],
       ["Pruebas automatizadas en verde", "100 % (hoy 172)", "Garantiza que las reglas del negocio se cumplen"],
       ["Disponibilidad del servicio", "Mas del 99 %", "Un CRM caido es una venta perdida"]],
      fs=12, widths=[3.6, 3.0, 5.3])
txt(sl, 0.7, 5.6, 11.9, 0.8,
    "Ninguna de estas metricas es de vanidad: todas se miden dentro del sistema y ninguna depende de una opinion.",
    size=14, bold=True, color=BLUE)

# ---------------------------------------------------------------- 18 cierre
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, 7.5, NAVY, shape=MSO_SHAPE.RECTANGLE)
rect(sl, 0, 0, 0.35, 7.5, BLUE, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 1.2, 2.3, 11, 1.4,
    "El problema de LYC no es conseguir clientes:\nes no perderlos por desorden.",
    size=32, bold=True, color=WHITE, space=10)
rect(sl, 1.2, 4.0, 3.0, 0.05, BLUE, shape=MSO_SHAPE.RECTANGLE)
txt(sl, 1.2, 4.35, 11, 1.2,
    ["Enlyce convierte ese desorden en procesos, reglas y responsables.",
     "Y lo hace por menos de lo que cuesta perder un solo negocio."],
    size=17, color=RGBColor(0xD8, 0xE4, 0xF2), space=8)
txt(sl, 1.2, 6.3, 11, 0.4, "Gracias.  Preguntas.", size=20, bold=True, color=RGBColor(0x7F, 0xB3, 0xFF))

prs.save(OUT)
print("ok", OUT, len(prs.slides.__iter__.__self__._sldIdLst))
