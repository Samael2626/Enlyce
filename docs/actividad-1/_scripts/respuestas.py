# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from acentos import fix

OUT = r"D:/Proyectos/Enlyce/docs/actividad-1/1_MODELO_DEL_NEGOCIO_RESPUESTAS.docx"
d = docx.Document()
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.2)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

def H(t, lvl=1):
    p = d.add_heading(fix(t), lvl)
    for r in p.runs: r.font.color.rgb = RGBColor(0x0F, 0x25, 0x40)
    return p

def P(t, bold=False):
    p = d.add_paragraph(); r = p.add_run(fix(t)); r.bold = bold; return p

def B(t): d.add_paragraph(fix(t), style="List Bullet")
def N(t): d.add_paragraph(fix(t), style="List Number")

def TAB(headers, rows):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        run = c.paragraphs[0].add_run(fix(h)); run.bold = True; run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(fix(v)); run.font.size = Pt(10)
    d.add_paragraph()

ti = d.add_heading("MODELO DEL NEGOCIO - ENLYCE", 0)
ti.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = P("CRM inmobiliario para L&C Propiedad Raiz (LYC)  |  Actividad 1  |  Ingenieria de Software  |  Agosto de 2026")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph()

H("0. Contexto del negocio modelado", 1)
P("L&C Propiedad Raiz es una inmobiliaria independiente de Medellin con cuatro asesores comerciales y menos "
  "de cincuenta inmuebles activos. Hoy trabaja con Wasi en plan gratuito (un solo usuario), hojas de calculo y "
  "conversaciones de WhatsApp sin ningun orden. El resultado es conocido: leads que nadie atiende, dos asesores "
  "llamando al mismo cliente y cero evidencia de cumplimiento de la Ley 1581 de 2012.")
P("Enlyce es el sistema que modela y automatiza ese negocio: captacion de leads, seguimiento por etapas, "
  "asignacion por asesor y cumplimiento normativo, en una sola plataforma. A continuacion se responden los cinco "
  "componentes del modelado de negocio aplicados a este caso real.")

H("1. Procesos de negocio", 1)
P("Un proceso es la secuencia de actividades que genera valor. Enlyce modela cuatro procesos principales.")

H("1.1 Proceso central: captacion y cierre de un lead inmobiliario", 2)
N("Un interesado contacta a LYC por WhatsApp, por el portal o por referido.")
N("El asesor registra el lead en Enlyce y captura la autorizacion de tratamiento de datos (requisito obligatorio: sin ella el sistema no permite crear el lead).")
N("El sistema asigna el lead a un asesor responsable, de forma manual o automatica.")
N("El lead entra al pipeline en la etapa Contacto y queda visible en el tablero Kanban.")
N("El asesor registra cada interaccion (llamada, mensaje, correo) sobre el lead.")
N("Se agenda una visita al inmueble y el lead avanza a la etapa Visita.")
N("Se presenta la propuesta economica: el lead pasa a Propuesta y luego a Negociacion.")
N("El negocio se cierra (Cerrado ganado) o se descarta con motivo (Cerrado perdido).")
N("El sistema deja registro de auditoria de todo el recorrido, para efectos legales y de productividad.")

H("1.2 Procesos de apoyo", 2)
TAB(["Proceso", "Que hace", "Valor que genera"], [
    ["Gestion de inmuebles", "Alta, actualizacion y baja logica de propiedades con su propietario asociado.",
     "Catalogo confiable y unico para todo el equipo."],
    ["Control de seguimiento", "Detecta leads sin gestion durante N dias y genera alertas al asesor.",
     "Evita que se pierda un lead por olvido."],
    ["Cumplimiento Ley 1581", "Consentimiento, politica de tratamiento versionada, atencion de derechos del titular y auditoria.",
     "Evita sanciones de la SIC y da confianza al cliente."],
    ["Administracion de usuarios", "Registro de asesores, autenticacion y control de roles.",
     "Cada quien ve y hace solo lo que le corresponde."],
])

H("2. Actividades del negocio", 1)
P("Las actividades son las tareas concretas dentro de cada proceso. En Enlyce se agrupan por modulo:")
TAB(["Modulo", "Actividades"], [
    ["Leads", "Registrar lead, capturar consentimiento, consultar, actualizar datos, dar de baja logica."],
    ["Pipeline", "Consultar el tablero, mover el lead de etapa, registrar la transicion, consultar el historial de etapas."],
    ["Asignacion", "Asignar lead a un asesor, reasignar, consultar la carga de trabajo por asesor."],
    ["Interacciones", "Registrar llamada, mensaje, correo o nota; consultar el historial del lead."],
    ["Visitas", "Agendar visita, confirmar, reprogramar, marcar como realizada o cancelada."],
    ["Alertas", "Detectar leads sin seguimiento, notificar al asesor responsable, marcar la alerta como atendida."],
    ["Inmuebles", "Publicar inmueble, asociar propietario, actualizar precio y estado, retirar de la vitrina."],
    ["Datos personales", "Otorgar y revocar consentimiento, publicar la politica vigente, atender solicitudes del titular, consultar la auditoria."],
    ["Seguridad", "Iniciar sesion, cerrar sesion, validar rol, registrar el evento en la auditoria."],
])

H("3. Reglas del negocio", 1)
P("Son las normas que el sistema obliga a cumplir. Se dividen en reglas legales, comerciales y de integridad.")

H("3.1 Reglas legales (Ley 1581 de 2012)", 2)
B("Ningun lead puede crearse sin autorizacion previa, expresa e informada del titular de los datos.")
B("La autorizacion se guarda con fecha, medio, finalidad y version de la politica aceptada.")
B("El titular puede revocar su autorizacion en cualquier momento; al revocar, sus datos dejan de usarse con fines comerciales.")
B("El titular tiene derecho a conocer, actualizar, rectificar y suprimir sus datos, y a solicitar prueba de la autorizacion.")
B("La politica de tratamiento es publica, versionada y debe estar vigente en el momento de la captura.")
B("Toda operacion sobre datos personales queda registrada en el log de auditoria.")

H("3.2 Reglas comerciales", 2)
B("Todo lead debe tener exactamente un asesor responsable: no existen leads sin dueno ni con dueno compartido.")
B("Un lead solo avanza por las etapas definidas del pipeline y no puede saltarse etapas hacia adelante.")
B("Un lead que llega a Cerrado perdido debe tener un motivo de descarte registrado.")
B("Un lead sin interaccion durante el plazo definido genera una alerta automatica al asesor responsable.")
B("Una visita solo se agenda sobre un inmueble disponible y con un lead activo.")

H("3.3 Reglas de integridad", 2)
B("Nada se elimina fisicamente: toda baja es logica y queda auditada.")
B("El correo electronico del asesor es unico dentro del sistema.")
B("Las contrasenas se almacenan siempre con hash (bcrypt), nunca en texto plano.")
B("Cada peticion se autentica por token firmado en cookie HttpOnly y se autoriza por rol.")
B("Un inmueble no se retira de la vitrina si tiene visitas agendadas pendientes.")

H("4. Actores del negocio", 1)
TAB(["Actor", "Tipo", "Rol en el negocio"], [
    ["Asesor comercial", "Humano - interno", "Registra y atiende leads, agenda visitas, mueve el pipeline. Usuario principal."],
    ["Administrador o gerente", "Humano - interno", "Gestiona asesores, revisa productividad, configura etapas y consulta auditoria."],
    ["Lead o cliente interesado", "Humano - externo", "Titular de los datos. Solicita informacion, visita inmuebles y compra o arrienda."],
    ["Propietario del inmueble", "Humano - externo", "Entrega el inmueble en gestion a la inmobiliaria."],
    ["Oficial de datos personales", "Humano - interno", "Responde solicitudes del titular y vela por la Ley 1581."],
    ["Superintendencia de Industria y Comercio", "Organizacion - externa", "Vigila el tratamiento de datos y exige el registro en el RNBD."],
    ["WhatsApp Business API (Meta)", "Sistema - externo", "Canal de entrada de leads y de comunicacion."],
    ["Servicio de correo transaccional", "Sistema - externo", "Entrega alertas y notificaciones."],
    ["Motor de alertas de Enlyce", "Sistema - interno", "Detecta leads sin seguimiento y notifica sin intervencion humana."],
])

H("5. Objetos del negocio", 1)
P("Son los elementos de informacion que circulan por los procesos. Corresponden a las entidades del dominio de Enlyce.")
TAB(["Objeto", "Atributos principales", "Para que sirve"], [
    ["Lead", "Nombre, telefono, correo, fuente, etapa actual, asesor asignado, estado, fecha de registro.",
     "Objeto central: representa la oportunidad comercial."],
    ["Asesor", "Nombre, correo, hash de contrasena, rol, estado activo.",
     "Identifica al responsable de cada lead y controla el acceso."],
    ["Inmueble", "Codigo, tipo, direccion, precio, area, estado, propietario.",
     "Producto que se ofrece al lead."],
    ["Propietario", "Nombre, documento, contacto.", "Dueno del inmueble en gestion."],
    ["Interaccion", "Lead, tipo (llamada, mensaje, correo, nota), fecha, descripcion, asesor.",
     "Historial de contacto: evidencia de la gestion comercial."],
    ["Visita", "Lead, inmueble, fecha y hora, estado (agendada, realizada, cancelada), observaciones.",
     "Momento clave de conversion del embudo."],
    ["Etapa del pipeline", "Nombre, orden, condiciones de avance.", "Define el flujo comercial estandar."],
    ["Transicion de pipeline", "Lead, etapa origen, etapa destino, fecha, asesor.",
     "Permite medir tiempos por etapa y tasa de conversion."],
    ["Consentimiento", "Titular, finalidad, medio, fecha, version de politica, estado (vigente o revocado).",
     "Prueba legal de la autorizacion exigida por la Ley 1581."],
    ["Politica de tratamiento", "Version, contenido, fecha de vigencia.", "Documento publico que respalda cada consentimiento."],
    ["Registro de auditoria", "Usuario, accion, entidad afectada, fecha, datos anteriores y nuevos.",
     "Trazabilidad completa: nada se pierde ni se borra en silencio."],
    ["Alerta de seguimiento", "Lead, asesor, motivo, fecha de generacion, estado.",
     "Convierte el olvido en una tarea visible."],
])

H("6. Conclusion", 1)
P("El modelado del negocio de LYC deja ver que el problema no es la falta de clientes, sino la falta de un proceso "
  "que los sostenga. Enlyce convierte ese proceso informal, disperso entre WhatsApp, hojas de calculo y memoria de "
  "los asesores, en procesos explicitos, actividades con responsable, reglas que el sistema hace cumplir, actores "
  "con permisos definidos y objetos de negocio trazables.")
P("El estado actual del proyecto respalda el modelo: la plataforma esta construida sobre arquitectura limpia "
  "(dominio sin dependencias externas), tiene implementados los modulos de leads, inmuebles, pipeline, interacciones, "
  "visitas, alertas, autenticacion y cumplimiento de la Ley 1581, y cuenta con 172 pruebas automatizadas en verde "
  "(133 de dominio y 39 de integracion).")

d.save(OUT)
print("ok", OUT)
